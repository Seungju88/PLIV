"""Qt GUI for the PLIV plugin."""

from __future__ import annotations

import json
import re
import sys
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any, List, Optional

from pymol.Qt import QtCore, QtGui, QtWidgets

from .config import ConfigManager
from .engine import InteractionEngine
from .models import AnalysisRequest, AnalysisSummary, LigandTarget
from .renderer import InteractionRenderer


class PLIVDialog(QtWidgets.QDialog):
    """Thin GUI wrapper that delegates logic to the engine and renderer."""

    def __init__(self, controller: "PLIVController") -> None:
        super().__init__()
        self.controller = controller
        self.setWindowTitle("PLIV - Protein-Ligand Interaction Visualizer")
        self.setGeometry(300, 220, 560, 760)
        self.setMinimumSize(540, 700)
        self._build_ui()

    def _build_ui(self) -> None:
        layout = QtWidgets.QVBoxLayout(self)
        layout.setSpacing(10)
        splitter = QtWidgets.QSplitter(QtCore.Qt.Vertical)
        splitter.setChildrenCollapsible(False)

        tab_widget = QtWidgets.QTabWidget()
        splitter.addWidget(tab_widget)

        def make_tab() -> tuple[QtWidgets.QWidget, QtWidgets.QVBoxLayout]:
            page = QtWidgets.QWidget()
            page_layout = QtWidgets.QVBoxLayout(page)
            page_layout.setContentsMargins(0, 0, 0, 0)
            page_layout.setSpacing(10)
            return page, page_layout

        def wrap_in_scroll(widget: QtWidgets.QWidget) -> QtWidgets.QScrollArea:
            scroll_area = QtWidgets.QScrollArea()
            scroll_area.setWidgetResizable(True)
            scroll_area.setFrameShape(QtWidgets.QFrame.NoFrame)
            scroll_area.setWidget(widget)
            return scroll_area

        analyze_page, analyze_layout = make_tab()
        session_page, session_layout = make_tab()
        view_page, view_layout = make_tab()

        selection_group = QtWidgets.QGroupBox("Selections")
        selection_layout = QtWidgets.QVBoxLayout(selection_group)

        mode_row = QtWidgets.QHBoxLayout()
        mode_row.addWidget(QtWidgets.QLabel("Selection Mode"))
        self.selection_mode_combo = QtWidgets.QComboBox()
        self.selection_mode_combo.addItem("Complex Object", "complex")
        self.selection_mode_combo.addItem("Docking Receptor + Ligand Objects", "docking")
        self.selection_mode_combo.currentIndexChanged.connect(self.controller.selection_mode_changed)
        mode_row.addWidget(self.selection_mode_combo, 1)
        selection_layout.addLayout(mode_row)

        self.primary_selection_label = QtWidgets.QLabel()
        selection_layout.addWidget(self.primary_selection_label)
        self.complex_list = QtWidgets.QListWidget()
        self.complex_list.setMinimumHeight(72)
        self.complex_list.setMaximumHeight(110)
        self.complex_list.itemClicked.connect(self._on_complex_selected)
        selection_layout.addWidget(self.complex_list)

        self.secondary_selection_label = QtWidgets.QLabel()
        selection_layout.addWidget(self.secondary_selection_label)
        self.ligand_list = QtWidgets.QListWidget()
        self.ligand_list.setSelectionMode(QtWidgets.QAbstractItemView.SingleSelection)
        self.ligand_list.setMinimumHeight(96)
        self.ligand_list.setMaximumHeight(132)
        selection_layout.addWidget(self.ligand_list)

        selection_tools_row = QtWidgets.QHBoxLayout()
        self.button_select_all_ligands = QtWidgets.QPushButton("Select All Ligands")
        self.button_select_all_ligands.clicked.connect(self.controller.select_all_ligands)
        selection_tools_row.addWidget(self.button_select_all_ligands)

        self.button_clear_ligand_selection = QtWidgets.QPushButton("Clear Ligand Selection")
        self.button_clear_ligand_selection.clicked.connect(self.controller.clear_ligand_selection)
        selection_tools_row.addWidget(self.button_clear_ligand_selection)

        selection_tools_row.addStretch(1)
        button_refresh = QtWidgets.QPushButton("Reload Object List")
        button_refresh.clicked.connect(self.controller.refresh_objects)
        selection_tools_row.addWidget(button_refresh)
        selection_layout.addLayout(selection_tools_row)
        analyze_layout.addWidget(selection_group)
        self._update_selection_labels()

        controls_group = QtWidgets.QGroupBox("Interaction Settings")
        controls_layout = QtWidgets.QVBoxLayout(controls_group)

        profile_row = QtWidgets.QHBoxLayout()
        profile_row.addWidget(QtWidgets.QLabel("Profile"))
        profile_value = QtWidgets.QLabel(self.controller.config.default_profile())
        profile_value.setStyleSheet(
            "background-color: #f3f4f6; border: 1px solid #d1d5db; "
            "border-radius: 6px; padding: 6px 10px; font-weight: 600;"
        )
        profile_row.addWidget(profile_value, 1)
        controls_layout.addLayout(profile_row)

        profile_tools_row = QtWidgets.QHBoxLayout()
        button_reload = QtWidgets.QPushButton("Reload Config")
        button_reload.clicked.connect(self.controller.reload_config)
        profile_tools_row.addStretch(1)
        profile_tools_row.addWidget(button_reload)
        controls_layout.addLayout(profile_tools_row)

        controls_layout.addWidget(QtWidgets.QLabel("Interaction Families"))
        interaction_grid = QtWidgets.QGridLayout()
        self.interaction_checks = {}
        self.interaction_color_buttons = {}
        for index, (key, label) in enumerate(self.controller.interaction_labels().items()):
            checkbox = QtWidgets.QCheckBox(label)
            checkbox.setChecked(bool(self.controller.config.get("interactions", key, "enabled", default=True)))
            self.interaction_checks[key] = checkbox
            color_button = QtWidgets.QPushButton()
            color_button.setFixedSize(26, 22)
            color_button.setToolTip(f"Choose color for {label}")
            color_button.clicked.connect(
                lambda _checked=False, interaction_key=key: self.controller.choose_interaction_color(interaction_key)
            )
            self.interaction_color_buttons[key] = color_button
            self.update_interaction_color_button(
                key,
                self.controller.interaction_color_display_value(key),
            )

            row_widget = QtWidgets.QWidget()
            row_layout = QtWidgets.QHBoxLayout(row_widget)
            row_layout.setContentsMargins(0, 0, 0, 0)
            row_layout.setSpacing(6)
            row_layout.addWidget(checkbox)
            row_layout.addWidget(color_button)
            row_layout.addStretch(1)
            interaction_grid.addWidget(row_widget, index // 2, index % 2)
        controls_layout.addLayout(interaction_grid)
        analyze_layout.addWidget(controls_group)

        save_image_group = QtWidgets.QGroupBox("Save Image")
        save_image_layout = QtWidgets.QVBoxLayout(save_image_group)
        save_image_hint = QtWidgets.QLabel(
            "Save the current PyMOL viewport as a PNG image for quick figure capture."
        )
        save_image_hint.setWordWrap(True)
        save_image_layout.addWidget(save_image_hint)

        button_save_image = QtWidgets.QPushButton("Save Image")
        button_save_image.clicked.connect(self.controller.save_image)
        save_image_layout.addWidget(button_save_image)
        session_layout.addWidget(save_image_group)

        session_group = QtWidgets.QGroupBox("Session Files")
        session_group_layout = QtWidgets.QVBoxLayout(session_group)
        session_hint = QtWidgets.QLabel(
            "Save or load a PyMOL session together with window size and analysis context."
        )
        session_hint.setWordWrap(True)
        session_group_layout.addWidget(session_hint)

        session_button_grid = QtWidgets.QGridLayout()
        button_save_session = QtWidgets.QPushButton("Save Session")
        button_save_session.clicked.connect(self.controller.save_session)
        session_button_grid.addWidget(button_save_session, 0, 0)

        button_load_session = QtWidgets.QPushButton("Load Session")
        button_load_session.clicked.connect(self.controller.load_session)
        session_button_grid.addWidget(button_load_session, 0, 1)
        session_group_layout.addLayout(session_button_grid)
        session_layout.addWidget(session_group)

        batch_group = QtWidgets.QGroupBox("Batch Export")
        batch_layout = QtWidgets.QVBoxLayout(batch_group)
        batch_hint = QtWidgets.QLabel(
            "Use saved views 1-5 to export ligand figures. In docking mode, batch export uses all ligand objects selected in the Analyze tab."
        )
        batch_hint.setWordWrap(True)
        batch_layout.addWidget(batch_hint)

        batch_views_row = QtWidgets.QHBoxLayout()
        batch_views_row.addWidget(QtWidgets.QLabel("Saved Views"))
        self.batch_view_checks = {}
        for slot in range(1, 6):
            checkbox = QtWidgets.QCheckBox(str(slot))
            self.batch_view_checks[slot] = checkbox
            batch_views_row.addWidget(checkbox)
        batch_views_row.addStretch(1)
        batch_layout.addLayout(batch_views_row)

        batch_style_row = QtWidgets.QHBoxLayout()
        batch_style_row.addWidget(QtWidgets.QLabel("Style"))
        self.batch_style_combo = QtWidgets.QComboBox()
        self.batch_style_combo.addItem("Current visible state", "current")
        self.batch_style_combo.addItem("Working view", "working")
        self.batch_style_combo.addItem("Publication style", "publication")
        batch_style_row.addWidget(self.batch_style_combo, 1)
        batch_layout.addLayout(batch_style_row)

        batch_render_grid = QtWidgets.QGridLayout()
        self.batch_ray_checkbox = QtWidgets.QCheckBox("Ray trace")
        self.batch_ray_checkbox.setChecked(True)
        batch_render_grid.addWidget(self.batch_ray_checkbox, 0, 0)
        batch_render_grid.addWidget(QtWidgets.QLabel("Width"), 0, 1)
        self.batch_width_spin = QtWidgets.QSpinBox()
        self.batch_width_spin.setRange(320, 7680)
        self.batch_width_spin.setSingleStep(160)
        self.batch_width_spin.setValue(1920)
        batch_render_grid.addWidget(self.batch_width_spin, 0, 2)
        batch_render_grid.addWidget(QtWidgets.QLabel("Height"), 0, 3)
        self.batch_height_spin = QtWidgets.QSpinBox()
        self.batch_height_spin.setRange(240, 4320)
        self.batch_height_spin.setSingleStep(120)
        self.batch_height_spin.setValue(1440)
        batch_render_grid.addWidget(self.batch_height_spin, 0, 4)
        batch_render_grid.addWidget(QtWidgets.QLabel("DPI"), 0, 5)
        self.batch_dpi_spin = QtWidgets.QSpinBox()
        self.batch_dpi_spin.setRange(72, 1200)
        self.batch_dpi_spin.setValue(300)
        batch_render_grid.addWidget(self.batch_dpi_spin, 0, 6)
        batch_layout.addLayout(batch_render_grid)

        batch_output_row = QtWidgets.QHBoxLayout()
        batch_output_row.addWidget(QtWidgets.QLabel("Output Folder"))
        self.batch_output_dir_edit = QtWidgets.QLineEdit()
        self.batch_output_dir_edit.setReadOnly(True)
        batch_output_row.addWidget(self.batch_output_dir_edit, 1)
        button_batch_output = QtWidgets.QPushButton("Choose Folder")
        button_batch_output.clicked.connect(self.controller.choose_batch_output_directory)
        batch_output_row.addWidget(button_batch_output)
        batch_layout.addLayout(batch_output_row)

        batch_action_row = QtWidgets.QHBoxLayout()
        button_batch_export = QtWidgets.QPushButton("Run Batch Export")
        button_batch_export.clicked.connect(self.controller.run_batch_export)
        batch_action_row.addWidget(button_batch_export)

        button_batch_docx = QtWidgets.QPushButton("Save DOCX Report")
        button_batch_docx.clicked.connect(self.controller.save_batch_docx_report)
        batch_action_row.addWidget(button_batch_docx)
        batch_layout.addLayout(batch_action_row)
        session_layout.addWidget(batch_group)

        action_group = QtWidgets.QGroupBox("Actions")
        action_layout = QtWidgets.QVBoxLayout(action_group)
        self.button_run = QtWidgets.QPushButton("Run Analysis")
        self.button_run.setStyleSheet(
            "background-color: #16a34a; color: white; font-weight: bold; min-height: 40px; "
            "border-radius: 8px;"
        )
        self.button_run.clicked.connect(self.controller.run_analysis)
        action_layout.addWidget(self.button_run)

        secondary_actions = QtWidgets.QHBoxLayout()
        button_publication = QtWidgets.QPushButton("Apply Publication Style")
        button_publication.setStyleSheet(
            "background-color: #0284c7; color: white; font-weight: bold; min-height: 36px; "
            "border-radius: 8px;"
        )
        button_publication.clicked.connect(self.controller.apply_publication_view)
        secondary_actions.addWidget(button_publication)

        button_revert = QtWidgets.QPushButton("Restore Working View")
        button_revert.setStyleSheet(
            "background-color: #64748b; color: white; font-weight: bold; min-height: 36px; "
            "border-radius: 8px;"
        )
        button_revert.clicked.connect(self.controller.revert_publication_view)
        secondary_actions.addWidget(button_revert)
        action_layout.addLayout(secondary_actions)

        button_clear = QtWidgets.QPushButton("Clear Plugin Objects")
        button_clear.clicked.connect(self.controller.clear_plugin_objects)
        action_layout.addWidget(button_clear)
        analyze_layout.addWidget(action_group)

        viewpoint_group = QtWidgets.QGroupBox("Saved Views")
        viewpoint_layout = QtWidgets.QGridLayout(viewpoint_group)
        for slot in range(1, 6):
            save_button = QtWidgets.QPushButton(f"Save View {slot}")
            save_button.clicked.connect(
                lambda _checked=False, selected_slot=slot: self.controller.store_viewpoint(selected_slot)
            )
            viewpoint_layout.addWidget(save_button, 0, slot - 1)

            load_button = QtWidgets.QPushButton(f"Load View {slot}")
            load_button.clicked.connect(
                lambda _checked=False, selected_slot=slot: self.controller.restore_viewpoint(selected_slot)
            )
            viewpoint_layout.addWidget(load_button, 1, slot - 1)
        view_hint = QtWidgets.QLabel(
            "Saved views restore only the camera position and zoom, not the full visualization state."
        )
        view_hint.setWordWrap(True)
        view_layout.addWidget(view_hint)
        view_layout.addWidget(viewpoint_group)

        analyze_layout.addStretch(1)
        session_layout.addStretch(1)
        view_layout.addStretch(1)

        tab_widget.addTab(wrap_in_scroll(analyze_page), "Analyze")
        tab_widget.addTab(wrap_in_scroll(session_page), "Session / Export")
        tab_widget.addTab(wrap_in_scroll(view_page), "View")

        status_group = QtWidgets.QGroupBox("Status Log")
        status_layout = QtWidgets.QVBoxLayout(status_group)
        self.status_box = QtWidgets.QPlainTextEdit()
        self.status_box.setReadOnly(True)
        self.status_box.setMinimumHeight(150)
        status_layout.addWidget(self.status_box)
        splitter.addWidget(status_group)
        splitter.setStretchFactor(0, 4)
        splitter.setStretchFactor(1, 2)
        splitter.setSizes([540, 220])
        layout.addWidget(splitter)

    def _on_complex_selected(self, item: QtWidgets.QListWidgetItem) -> None:
        self.controller.populate_ligands(item.text())

    def selected_input_mode(self) -> str:
        return str(self.selection_mode_combo.currentData() or "complex")

    def selected_complex(self) -> Optional[str]:
        item = self.complex_list.currentItem()
        return item.text() if item else None

    def selected_ligand(self) -> Optional[LigandTarget]:
        item = self.ligand_list.currentItem()
        if item is not None:
            payload = item.data(QtCore.Qt.UserRole)
            if isinstance(payload, LigandTarget):
                return payload
            chain, resi, resn = item.text().split(":")
            return LigandTarget(chain=chain, resi=resi, resn=resn)

        ligands = self.selected_ligands_for_batch()
        return ligands[0] if ligands else None

    def selected_ligands_for_batch(self) -> List[LigandTarget]:
        items = self.ligand_list.selectedItems()
        if not items:
            current_item = self.ligand_list.currentItem()
            if current_item is not None:
                items = [current_item]

        ligands: List[LigandTarget] = []
        seen_labels = set()
        for item in items:
            payload = item.data(QtCore.Qt.UserRole)
            if isinstance(payload, LigandTarget):
                ligand = payload
            else:
                chain, resi, resn = item.text().split(":")
                ligand = LigandTarget(chain=chain, resi=resi, resn=resn)
            if ligand.label in seen_labels:
                continue
            ligands.append(ligand)
            seen_labels.add(ligand.label)
        return ligands

    def clear_ligand_selection(self) -> None:
        self.ligand_list.clearSelection()
        self.ligand_list.setCurrentRow(-1)

    def select_all_ligands(self) -> int:
        if self.ligand_list.selectionMode() == QtWidgets.QAbstractItemView.SingleSelection:
            if self.ligand_list.count():
                self.ligand_list.setCurrentRow(0)
                return 1
            return 0
        self.ligand_list.selectAll()
        return len(self.ligand_list.selectedItems())

    def selected_interaction_types(self) -> List[str]:
        enabled = []
        for key, checkbox in self.interaction_checks.items():
            if checkbox.isChecked():
                enabled.append(key)
        return enabled

    def selected_profile(self) -> str:
        return self.controller.config.default_profile()

    def set_complexes(self, names: List[str]) -> None:
        self.complex_list.clear()
        self.complex_list.addItems(names)
        self.ligand_list.clear()
        self._sync_ligand_selection_controls()

    def set_ligands(self, ligands: List[LigandTarget]) -> None:
        self.ligand_list.clear()
        for ligand in ligands:
            item = QtWidgets.QListWidgetItem(ligand.label)
            item.setData(QtCore.Qt.UserRole, ligand)
            self.ligand_list.addItem(item)
        self._sync_ligand_selection_controls()

    def select_complex(self, complex_name: str) -> bool:
        matches = self.complex_list.findItems(complex_name, QtCore.Qt.MatchExactly)
        if not matches:
            return False
        item = matches[0]
        self.complex_list.setCurrentItem(item)
        self._on_complex_selected(item)
        return True

    def select_ligand(self, ligand: LigandTarget) -> bool:
        matches = self.ligand_list.findItems(ligand.label, QtCore.Qt.MatchExactly)
        if not matches:
            return False
        self.clear_ligand_selection()
        self.ligand_list.setCurrentItem(matches[0])
        matches[0].setSelected(True)
        return True

    def set_interaction_types(self, enabled_types: List[str]) -> None:
        enabled = set(enabled_types)
        for key, checkbox in self.interaction_checks.items():
            checkbox.setChecked(key in enabled)

    def sync_interaction_controls(self) -> None:
        for key, checkbox in self.interaction_checks.items():
            checkbox.setChecked(bool(self.controller.config.get("interactions", key, "enabled", default=True)))
            self.update_interaction_color_button(key, self.controller.interaction_color_display_value(key))

    def set_selection_mode(self, mode: str) -> bool:
        index = self.selection_mode_combo.findData(mode)
        if index < 0:
            return False
        self.selection_mode_combo.setCurrentIndex(index)
        self._update_selection_labels()
        return True

    def _update_selection_labels(self) -> None:
        if self.selected_input_mode() == "docking":
            self.primary_selection_label.setText("Step 1: Select receptor object")
            self.secondary_selection_label.setText("Step 2: Select one or more ligand objects")
        else:
            self.primary_selection_label.setText("Step 1: Select complex object")
            self.secondary_selection_label.setText("Step 2: Select ligand instance (chain:resi:resn)")
        self._sync_ligand_selection_controls()

    def _sync_ligand_selection_controls(self) -> None:
        docking_mode = self.selected_input_mode() == "docking"
        selection_mode = (
            QtWidgets.QAbstractItemView.ExtendedSelection
            if docking_mode
            else QtWidgets.QAbstractItemView.SingleSelection
        )
        self.ligand_list.setSelectionMode(selection_mode)
        self.button_select_all_ligands.setEnabled(docking_mode and self.ligand_list.count() > 0)
        self.button_clear_ligand_selection.setEnabled(self.ligand_list.count() > 0)

    def update_interaction_color_button(self, interaction_type: str, color_value: str) -> None:
        button = self.interaction_color_buttons.get(interaction_type)
        if button is None:
            return
        color = QtGui.QColor(color_value)
        if not color.isValid():
            color = QtGui.QColor("#9ca3af")
        border = "#111827" if color.lightness() > 120 else "#e5e7eb"
        button.setStyleSheet(
            "QPushButton {"
            f"background-color: {color.name()};"
            f"border: 1px solid {border};"
            "border-radius: 5px;"
            "}"
        )

    def log(self, message: str) -> None:
        self.status_box.appendPlainText(message)

    def set_run_analysis_enabled(self, enabled: bool, reason: Optional[str] = None) -> None:
        self.button_run.setEnabled(enabled)
        self.button_run.setToolTip(reason or "")

    def batch_selected_view_slots(self) -> List[int]:
        return [slot for slot, checkbox in self.batch_view_checks.items() if checkbox.isChecked()]

    def batch_export_style(self) -> str:
        return str(self.batch_style_combo.currentData() or "current")

    def batch_export_ray(self) -> bool:
        return self.batch_ray_checkbox.isChecked()

    def batch_export_dimensions(self) -> tuple[int, int]:
        return int(self.batch_width_spin.value()), int(self.batch_height_spin.value())

    def batch_export_dpi(self) -> int:
        return int(self.batch_dpi_spin.value())

    def batch_output_directory(self) -> str:
        return self.batch_output_dir_edit.text().strip()

    def set_batch_output_directory(self, path: str) -> None:
        self.batch_output_dir_edit.setText(path)


class PLIVController:
    """Coordinates the GUI, engine, and renderer."""

    def __init__(self) -> None:
        self.config = ConfigManager()
        self.engine = InteractionEngine(self.config)
        self.renderer = InteractionRenderer(self.config)
        self.dialog = PLIVDialog(self)
        self._last_request: Optional[AnalysisRequest] = None
        self._last_summary = None
        self._publication_active = False
        self.refresh_objects()

        dependency_issue = self.engine.dependency_issue()
        if dependency_issue is not None:
            message = (
                dependency_issue
                + "\n\nPLIV requires a working NumPy installation in the same Python interpreter that PyMOL is using."
            )
            self.dialog.set_run_analysis_enabled(False, message)
            self.dialog.log(message)
            QtWidgets.QMessageBox.warning(self.dialog, "PLIV Dependency Issue", message)

        self.dialog.set_batch_output_directory(str(self._default_directory()))

    def show(self) -> None:
        self.dialog.show()
        self.dialog.raise_()
        self.dialog.activateWindow()

    def interaction_labels(self):
        return {
            "hydrogen_bond": "Hydrogen bond",
            "halogen_bond": "Halogen bond",
            "salt_bridge": "Salt bridge",
            "pi_pi_stacking": "Pi-pi stacking",
            "pi_cation": "Pi-cation",
            "hydrophobic_contact": "Hydrophobic contact",
        }

    def interaction_color_display_value(self, interaction_type: str) -> str:
        raw_value = str(self.config.get("interactions", interaction_type, "color", default="#9ca3af"))
        qcolor = QtGui.QColor(raw_value)
        if qcolor.isValid():
            return qcolor.name()

        try:
            red, green, blue = self.engine.cmd.get_color_tuple(raw_value)
            return QtGui.QColor.fromRgbF(float(red), float(green), float(blue)).name()
        except Exception:
            return "#9ca3af"

    def choose_interaction_color(self, interaction_type: str) -> None:
        label = self.interaction_labels().get(interaction_type, interaction_type)
        initial_color = QtGui.QColor(self.interaction_color_display_value(interaction_type))
        selected = QtWidgets.QColorDialog.getColor(
            initial_color,
            self.dialog,
            f"Choose color for {label}",
        )
        if not selected.isValid():
            return

        color_value = selected.name()
        self.config.set_interaction_color(interaction_type, color_value)
        self.dialog.update_interaction_color_button(interaction_type, color_value)
        self._update_cached_interaction_color(interaction_type, color_value)
        if self._last_request is not None:
            self.renderer.update_interaction_family_color(self._last_request, interaction_type, color_value)
        self.dialog.log(f"Updated {label} color to {color_value}.")

    def _update_cached_interaction_color(self, interaction_type: str, color_value: str) -> None:
        if self._last_summary is None:
            return
        interactions = getattr(self._last_summary, "interactions", None)
        if not isinstance(interactions, list):
            return
        for record in interactions:
            if isinstance(record, dict) and record.get("kind") == interaction_type:
                record["color"] = color_value

    def selection_mode_changed(self) -> None:
        self.dialog._update_selection_labels()
        self.refresh_objects()
        mode = self.dialog.selected_input_mode()
        if mode == "docking":
            self.dialog.log("Switched selection mode to docking receptor + ligand objects.")
        else:
            self.dialog.log("Switched selection mode to complex object + ligand instance.")

    def refresh_objects(self) -> None:
        mode = self.dialog.selected_input_mode()
        if mode == "docking":
            names = self.engine.list_receptor_objects()
            noun = "receptor object"
        else:
            names = self.engine.list_receptor_objects()
            noun = "protein-containing object"
        self.dialog.set_complexes(names)
        self.dialog.log(f"Loaded {len(names)} {noun}(s).")

    def populate_ligands(self, complex_name: str) -> None:
        mode = self.dialog.selected_input_mode()
        if mode == "docking":
            ligands = self.engine.list_docking_ligand_objects(complex_name)
            label = "ligand object"
        else:
            ligands = self.engine.list_ligands(complex_name)
            label = "ligand instance"
        self.dialog.set_ligands(ligands)
        self.dialog.log(f"Loaded {len(ligands)} {label}(s) from {complex_name}.")

    def reload_config(self) -> None:
        self.config.reload()
        self.dialog.sync_interaction_controls()
        self.dialog.log(f"Reloaded config from {self.config.path}.")

    def clear_plugin_objects(self) -> None:
        self.renderer.clear_plugin_objects()
        self._publication_active = False
        self.dialog.log("Cleared plugin-owned objects.")

    def run_analysis(self) -> None:
        complex_name = self.dialog.selected_complex()
        ligand = self.dialog.selected_ligand()
        interaction_types = self.dialog.selected_interaction_types()
        profile_name = self.dialog.selected_profile()
        input_mode = self.dialog.selected_input_mode()

        if not complex_name or not ligand:
            if input_mode == "docking":
                self.dialog.log("Select both a receptor object and a ligand object first.")
            else:
                self.dialog.log("Select both a complex object and a ligand instance first.")
            return

        if not interaction_types:
            self.dialog.log("Enable at least one interaction family.")
            return

        if input_mode == "complex" and int(self.engine.cmd.count_atoms(f"({complex_name}) and polymer.protein")) == 0:
            self.dialog.log(
                "The selected object does not contain protein atoms. If the ligand was loaded as a separate PDB object, use Docking Receptor + Ligand Objects mode."
            )
            return

        self.engine.cmd.set("valence", 0)
        self.renderer.clear_plugin_objects()
        request = self.engine.build_request(
            complex_name=complex_name,
            ligand=ligand,
            interaction_types=interaction_types,
            profile_name=profile_name,
            input_mode=input_mode,
        )
        summary = self.engine.run_analysis(request)
        self.renderer.render_working_view(request, summary)
        self._last_request = request
        self._last_summary = summary
        self._publication_active = False

        self.dialog.log(
            f"Prepared analysis for {ligand.label} with {len(interaction_types)} interaction family(s)."
        )
        self.dialog.log(
            f"Pocket atoms: {summary.pocket_atom_count}, ligand atoms: {summary.ligand_atom_count}."
        )
        water_atom_count = getattr(summary, "water_atom_count", 0)
        if water_atom_count:
            self.dialog.log(f"Nearby water atoms: {water_atom_count}")
        ion_atom_count = getattr(summary, "ion_atom_count", 0)
        if ion_atom_count:
            self.dialog.log(f"Nearby ion atoms: {ion_atom_count}")
        interaction_counts = getattr(summary, "interaction_counts", {})
        for kind, count in sorted(interaction_counts.items()):
            self.dialog.log(f"* {kind}: {count}")
        for note in summary.notes:
            self.dialog.log(f"- {note}")

    def apply_publication_view(self) -> None:
        if self._last_request is None:
            self.dialog.log("Run an analysis first so the publication preset has a target.")
            return

        if not self._publication_active:
            self.renderer.store_view_state(self._last_request)
        self.renderer.apply_publication_view(self._last_request, self._last_summary)
        self._publication_active = True
        self.dialog.log("Applied publication view preset.")

    def revert_publication_view(self) -> None:
        if self._last_request is None:
            self.dialog.log("Run an analysis first so there is a view to restore.")
            return

        restored = self.renderer.restore_view_state(self._last_request)
        if restored:
            self._publication_active = False
            self.dialog.log("Reverted to the pre-publication view.")
        else:
            self.dialog.log("No saved pre-publication view was available.")

    def select_all_ligands(self) -> None:
        count = self.dialog.select_all_ligands()
        if self.dialog.selected_input_mode() != "docking":
            self.dialog.log("Ligand multi-selection is available in docking mode.")
            return
        self.dialog.log(f"Selected {count} ligand object(s) for batch export.")

    def clear_ligand_selection(self) -> None:
        self.dialog.clear_ligand_selection()
        self.dialog.log("Cleared ligand selection.")

    def choose_batch_output_directory(self) -> None:
        start_dir = self.dialog.batch_output_directory() or str(self._default_directory())
        directory = QtWidgets.QFileDialog.getExistingDirectory(
            self.dialog,
            "Choose Batch Export Folder",
            str(start_dir),
        )
        if not directory:
            return
        self.dialog.set_batch_output_directory(directory)
        self.dialog.log(f"Batch export folder set to {directory}.")

    def _collect_batch_export_settings(self) -> Optional[dict[str, Any]]:
        view_slots = self.dialog.batch_selected_view_slots()
        if not view_slots:
            self.dialog.log("Select at least one saved view slot for batch export.")
            return None

        input_mode = self.dialog.selected_input_mode()
        complex_name = self.dialog.selected_complex()
        if not complex_name:
            if input_mode == "docking":
                self.dialog.log("Select a receptor object before running batch export.")
            else:
                self.dialog.log("Select a complex object before running batch export.")
            return None

        ligands = self.dialog.selected_ligands_for_batch()
        if not ligands:
            if input_mode == "docking":
                self.dialog.log("Select one or more ligand objects for batch export.")
            else:
                self.dialog.log("Select a ligand instance for batch export.")
            return None

        style = self.dialog.batch_export_style()
        if style == "current" and len(ligands) > 1:
            self.dialog.log(
                "Current visible state can export only one ligand target at a time. Use Working view or Publication style for multi-ligand export."
            )
            return None

        interaction_types = self.dialog.selected_interaction_types()
        if style in {"working", "publication"} and not interaction_types:
            self.dialog.log("Enable at least one interaction family before running working or publication batch export.")
            return None

        if (
            style in {"working", "publication"}
            and input_mode == "complex"
            and int(self.engine.cmd.count_atoms(f"({complex_name}) and polymer.protein")) == 0
        ):
            self.dialog.log(
                "The selected object does not contain protein atoms. If the ligand was loaded as a separate PDB object, use Docking Receptor + Ligand Objects mode."
            )
            return None

        output_dir_text = self.dialog.batch_output_directory() or str(self._default_directory())
        output_dir = Path(output_dir_text).expanduser()
        output_dir.mkdir(parents=True, exist_ok=True)
        self.dialog.set_batch_output_directory(str(output_dir))

        width, height = self.dialog.batch_export_dimensions()
        batch_token = self._batch_base_token(complex_name, ligands)
        return {
            "view_slots": view_slots,
            "style": style,
            "output_dir": output_dir,
            "width": width,
            "height": height,
            "dpi": self.dialog.batch_export_dpi(),
            "ray": self.dialog.batch_export_ray(),
            "input_mode": input_mode,
            "complex_name": complex_name,
            "ligands": ligands,
            "interaction_types": interaction_types,
            "profile_name": self.dialog.selected_profile(),
            "batch_token": batch_token,
        }

    def _build_batch_request(self, settings: dict[str, Any], ligand: LigandTarget) -> AnalysisRequest:
        return self.engine.build_request(
            complex_name=str(settings["complex_name"]),
            ligand=ligand,
            interaction_types=list(settings["interaction_types"]),
            profile_name=str(settings["profile_name"]),
            input_mode=str(settings["input_mode"]),
        )

    def _build_batch_target_manifest_entry(
        self,
        request: AnalysisRequest,
        summary: Optional[AnalysisSummary],
        *,
        status: str = "ready",
        error: str = "",
    ) -> dict[str, Any]:
        interaction_counts = {}
        notes: list[str] = []
        if summary is not None:
            interaction_counts = dict(getattr(summary, "interaction_counts", {}) or {})
            notes = [str(note) for note in getattr(summary, "notes", [])]

        payload = {
            "target_token": request.token,
            "ligand_label": request.ligand.label,
            "complex_name": request.complex_name,
            "input_mode": request.input_mode,
            "interaction_types": list(request.interaction_types),
            "interaction_counts": interaction_counts,
            "notes": notes,
            "status": status,
        }
        if error:
            payload["error"] = error
        return payload

    def _export_saved_views_for_target(
        self,
        request: AnalysisRequest,
        summary: Optional[AnalysisSummary],
        settings: dict[str, Any],
        timestamp: str,
        manifest_entries: list[dict[str, Any]],
    ) -> int:
        style = str(settings["style"])
        output_dir = Path(settings["output_dir"])
        width = int(settings["width"])
        height = int(settings["height"])
        dpi = int(settings["dpi"])
        ray = bool(settings["ray"])
        ray_flag = 1 if ray else 0
        exported_count = 0
        interaction_counts = {}
        if summary is not None:
            interaction_counts = dict(getattr(summary, "interaction_counts", {}) or {})

        for slot in list(settings["view_slots"]):
            restored = self.renderer.restore_viewpoint(slot)
            if not restored:
                self.dialog.log(f"Skipped {request.ligand.label} view {slot}: no saved viewpoint was available.")
                manifest_entries.append(
                    {
                        "target_token": request.token,
                        "ligand_label": request.ligand.label,
                        "complex_name": request.complex_name,
                        "input_mode": request.input_mode,
                        "view_slot": slot,
                        "style": style,
                        "status": "missing_viewpoint",
                    }
                )
                continue

            output_path = output_dir / self._suggested_batch_image_name(request.token, slot, style, timestamp)
            self.engine.cmd.png(
                str(output_path),
                width=width,
                height=height,
                dpi=dpi,
                ray=ray_flag,
                quiet=0,
            )
            exported_count += 1
            manifest_entries.append(
                {
                    "target_token": request.token,
                    "ligand_label": request.ligand.label,
                    "complex_name": request.complex_name,
                    "input_mode": request.input_mode,
                    "interaction_counts": interaction_counts,
                    "view_slot": slot,
                    "style": style,
                    "ray": ray,
                    "width": width,
                    "height": height,
                    "dpi": dpi,
                    "status": "saved",
                    "file": str(output_path),
                }
            )
            self.dialog.log(f"Batch-saved {request.ligand.label} view {slot} to {output_path}.")

        return exported_count

    def _execute_batch_export(
        self,
        settings: dict[str, Any],
        *,
        timestamp: Optional[str] = None,
    ) -> tuple[str, Path, dict[str, Any], int]:
        view_slots = list(settings["view_slots"])
        style = str(settings["style"])
        output_dir = Path(settings["output_dir"])
        ray = bool(settings["ray"])
        ligands = list(settings["ligands"])
        timestamp = timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")
        scene_name = f"mp_scene_batch_export_{timestamp}"
        manifest_entries: list[dict[str, Any]] = []
        target_entries: list[dict[str, Any]] = []
        exported_count = 0
        original_request = self._last_request
        original_summary = self._last_summary
        original_publication_active = self._publication_active

        self.renderer.store_scene(scene_name)
        try:
            if style == "current":
                request = self._build_batch_request(settings, ligands[0])
                summary = None
                if self._last_request is not None and self._last_summary is not None and self._last_request.token == request.token:
                    summary = self._last_summary
                target_entry = self._build_batch_target_manifest_entry(request, summary, status="rendered")
                saved_for_target = self._export_saved_views_for_target(
                    request,
                    summary,
                    settings,
                    timestamp,
                    manifest_entries,
                )
                target_entry["saved_images"] = saved_for_target
                exported_count += saved_for_target
                target_entries.append(target_entry)
            else:
                self.engine.cmd.set("valence", 0)
                for ligand in ligands:
                    request = self._build_batch_request(settings, ligand)
                    try:
                        self.renderer.clear_plugin_objects()
                        summary = self.engine.run_analysis(request)
                        self.renderer.render_working_view(request, summary)
                        if style == "publication":
                            self.renderer.apply_publication_view(request, summary)
                    except Exception as exc:
                        error_message = str(exc)
                        self.dialog.log(f"Skipped {ligand.label}: {error_message}")
                        target_entries.append(
                            self._build_batch_target_manifest_entry(
                                request,
                                None,
                                status="analysis_failed",
                                error=error_message,
                            )
                        )
                        continue

                    self._last_request = request
                    self._last_summary = summary
                    self._publication_active = style == "publication"
                    target_entry = self._build_batch_target_manifest_entry(request, summary, status="rendered")
                    saved_for_target = self._export_saved_views_for_target(
                        request,
                        summary,
                        settings,
                        timestamp,
                        manifest_entries,
                    )
                    target_entry["saved_images"] = saved_for_target
                    exported_count += saved_for_target
                    target_entries.append(target_entry)
        finally:
            restored_scene = self.renderer.restore_scene(scene_name)
            self.renderer.clear_scene(scene_name)
            self._last_request = original_request
            self._last_summary = original_summary
            self._publication_active = original_publication_active
            if restored_scene:
                self.dialog.log("Restored the pre-export visualization state.")

        manifest_path = output_dir / self._suggested_batch_manifest_name(str(settings["batch_token"]), timestamp)
        manifest_payload = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "style": style,
            "ray": ray,
            "width": int(settings["width"]),
            "height": int(settings["height"]),
            "dpi": int(settings["dpi"]),
            "views": list(view_slots),
            "input_mode": str(settings["input_mode"]),
            "complex_name": str(settings["complex_name"]),
            "batch_token": str(settings["batch_token"]),
            "target_count": len(ligands),
            "targets": target_entries,
            "images": manifest_entries,
        }
        manifest_path.write_text(json.dumps(manifest_payload, indent=2), encoding="utf-8")
        self.dialog.log(f"Saved batch export manifest to {manifest_path}.")
        return timestamp, manifest_path, manifest_payload, exported_count

    def run_batch_export(self) -> None:
        settings = self._collect_batch_export_settings()
        if settings is None:
            return
        _timestamp, _manifest_path, manifest_payload, exported_count = self._execute_batch_export(settings)
        self.dialog.log(
            f"Batch export complete: {exported_count} image(s) across {manifest_payload['target_count']} ligand target(s)."
        )

    def _load_docx_dependencies(self):
        try:
            from docx import Document
            from docx.shared import Inches
        except Exception as exc:
            message = (
                "DOCX report export requires python-docx in the same Python interpreter that PyMOL is using. "
                f"Interpreter: {sys.executable}. Original error: {exc}"
            )
            self.dialog.log(message)
            QtWidgets.QMessageBox.warning(self.dialog, "PLIV DOCX Dependency Issue", message)
            return None
        return Document, Inches

    def save_batch_docx_report(self) -> None:
        dependencies = self._load_docx_dependencies()
        if dependencies is None:
            return

        settings = self._collect_batch_export_settings()
        if settings is None:
            return

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_path = Path(settings["output_dir"]) / self._suggested_batch_report_name(str(settings["batch_token"]), timestamp)
        file_path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self.dialog,
            "Save DOCX Report",
            str(default_path),
            "Word Documents (*.docx)",
        )
        if not file_path:
            return

        report_path = Path(file_path)
        if report_path.suffix.lower() != ".docx":
            report_path = report_path.with_suffix(".docx")

        _timestamp, manifest_path, manifest_payload, exported_count = self._execute_batch_export(
            settings,
            timestamp=timestamp,
        )
        if exported_count == 0:
            self.dialog.log("No images were exported, so the DOCX report was not created.")
            return

        try:
            Document, Inches = dependencies
            document = Document()
            document.add_heading("PLIV Batch Export Report", 0)
            document.add_paragraph(f"Generated: {manifest_payload['generated_at']}")
            document.add_paragraph(f"Receptor / Complex: {manifest_payload['complex_name']}")
            document.add_paragraph(f"Selection mode: {manifest_payload['input_mode']}")
            document.add_paragraph(f"Ligand targets: {manifest_payload['target_count']}")
            document.add_paragraph(
                f"Style: {manifest_payload['style']} | Ray trace: {manifest_payload['ray']} | "
                f"Size: {manifest_payload['width']}x{manifest_payload['height']} | DPI: {manifest_payload['dpi']}"
            )
            document.add_paragraph(f"Manifest: {manifest_path}")

            saved_entries = [entry for entry in manifest_payload.get("images", []) if entry.get("status") == "saved"]
            entries_by_target: dict[str, list[dict[str, Any]]] = {}
            for entry in saved_entries:
                entries_by_target.setdefault(str(entry.get("target_token", "")), []).append(entry)

            target_entries = list(manifest_payload.get("targets", []))
            for target_index, target_entry in enumerate(target_entries, start=1):
                target_token = str(target_entry.get("target_token", ""))
                ligand_label = str(target_entry.get("ligand_label", target_token or f"Target {target_index}"))
                document.add_heading(ligand_label, level=1)
                document.add_paragraph(f"Target token: {target_token}")

                interaction_types = list(target_entry.get("interaction_types", []))
                if interaction_types:
                    document.add_paragraph("Interaction families: " + ", ".join(interaction_types))

                interaction_counts = dict(target_entry.get("interaction_counts", {}) or {})
                if interaction_counts:
                    counts_text = ", ".join(
                        f"{kind}={count}" for kind, count in sorted(interaction_counts.items())
                    )
                    document.add_paragraph("Interaction counts: " + counts_text)

                status = str(target_entry.get("status", ""))
                if status and status != "rendered":
                    document.add_paragraph(f"Status: {status}")
                error_message = str(target_entry.get("error", "")).strip()
                if error_message:
                    document.add_paragraph(f"Error: {error_message}")

                notes = [str(note) for note in target_entry.get("notes", [])]
                if notes:
                    document.add_paragraph("Analysis Notes")
                    for note in notes:
                        document.add_paragraph(note, style="List Bullet")

                target_images = sorted(
                    entries_by_target.get(target_token, []),
                    key=lambda entry: int(entry.get("view_slot", 0)),
                )
                if not target_images:
                    document.add_paragraph("No images were saved for this target.")
                else:
                    for entry in target_images:
                        document.add_heading(f"View {entry['view_slot']}", level=2)
                        document.add_paragraph(f"Image file: {entry['file']}")
                        document.add_picture(str(entry['file']), width=Inches(6.5))

                if target_index < len(target_entries):
                    document.add_page_break()

            document.save(str(report_path))
        except Exception as exc:
            self.dialog.log(f"Could not build the DOCX report: {exc}")
            QtWidgets.QMessageBox.warning(self.dialog, "PLIV DOCX Export Error", str(exc))
            return

        self.dialog.log(
            f"Batch export complete: {exported_count} image(s) across {manifest_payload['target_count']} ligand target(s)."
        )
        self.dialog.log(f"Saved DOCX report to {report_path}.")

    def save_image(self) -> None:
        default_path = self._default_directory() / self._suggested_image_name()
        file_path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self.dialog,
            "Save Image",
            str(default_path),
            "PNG Files (*.png)",
        )
        if not file_path:
            return

        output_path = Path(file_path)
        if output_path.suffix.lower() != ".png":
            output_path = output_path.with_suffix(".png")

        self.engine.cmd.png(str(output_path), dpi=300, ray=0, quiet=0)
        self.dialog.log(f"Saved image to {output_path}.")

    def save_session(self) -> None:
        default_path = self._default_directory() / self._suggested_session_name()
        file_path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self.dialog,
            "Save Session",
            str(default_path),
            "PyMOL Session Files (*.pse)",
        )
        if not file_path:
            return

        session_path = Path(file_path)
        if session_path.suffix.lower() != ".pse":
            session_path = session_path.with_suffix(".pse")

        self.engine.cmd.save(str(session_path), format="pse", quiet=0)
        self.dialog.set_batch_output_directory(str(session_path.parent))
        metadata = {
            "saved_at": datetime.now().isoformat(timespec="seconds"),
            "window": self._capture_pymol_window_state(),
            "analysis": self._capture_analysis_state(),
            "viewpoints": self.renderer.serialize_viewpoints(),
        }
        metadata_path = self._session_metadata_path(session_path)
        metadata_path.write_text(json.dumps(metadata, indent=2), encoding="utf-8")
        self.dialog.log(f"Saved session to {session_path}.")
        self.dialog.log(f"Saved session metadata to {metadata_path}.")

    def load_session(self) -> None:
        default_path = self._default_directory()
        file_path, _ = QtWidgets.QFileDialog.getOpenFileName(
            self.dialog,
            "Load Session",
            str(default_path),
            "PyMOL Session Files (*.pse)",
        )
        if not file_path:
            return

        session_path = Path(file_path)
        self.engine.cmd.load(str(session_path), partial=0, quiet=0)
        self._last_request = None
        self._last_summary = None
        self._publication_active = False
        self.renderer.clear_viewpoints()
        self.refresh_objects()
        self.dialog.set_batch_output_directory(str(session_path.parent))
        self.dialog.log(f"Loaded session from {session_path}.")

        metadata_path = self._session_metadata_path(session_path)
        if metadata_path.exists():
            try:
                metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
            except Exception as exc:
                self.dialog.log(f"Could not read session metadata: {exc}")
                return

            restored_viewpoints = self.renderer.load_serialized_viewpoints(metadata.get("viewpoints"))
            if restored_viewpoints:
                self.dialog.log(f"Restored {restored_viewpoints} saved view(s) from session metadata.")
            else:
                self.dialog.log("No saved view metadata was available in session metadata.")

            restored_analysis = self._restore_analysis_state(metadata.get("analysis"))
            if restored_analysis:
                self.dialog.log("Restored saved analysis context for publication tools.")
            else:
                self.dialog.log("No saved analysis context was available in session metadata.")
            QtCore.QTimer.singleShot(
                150,
                lambda metadata=metadata: self._restore_pymol_window_state(metadata.get("window", {})),
            )
            self.dialog.log(f"Queued window-size restore from {metadata_path}.")
        else:
            self.dialog.log("No session metadata sidecar was found for analysis or window restore.")

    def _default_directory(self) -> Path:
        session_file = str(self.engine.cmd.get("session_file") or "").strip()
        if session_file:
            try:
                return Path(session_file).expanduser().resolve().parent
            except Exception:
                pass
        return Path.home()

    def _suggested_image_name(self) -> str:
        token = self._last_request.token if self._last_request is not None else "pliv"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"{token}_{timestamp}.png"

    def _suggested_session_name(self) -> str:
        token = self._last_request.token if self._last_request is not None else "pliv_session"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"{token}_{timestamp}.pse"

    def _suggested_batch_image_name(self, token: str, slot: int, style: str, timestamp: str) -> str:
        return f"{token}_{style}_view{slot}_{timestamp}.png"

    def _suggested_batch_manifest_name(self, batch_token: str, timestamp: str) -> str:
        return f"{batch_token}_batch_export_{timestamp}.json"

    def _suggested_batch_report_name(self, batch_token: str, timestamp: str) -> str:
        return f"{batch_token}_batch_report_{timestamp}.docx"

    def _sanitize_token(self, value: str) -> str:
        cleaned = re.sub(r"[^A-Za-z0-9_]+", "_", str(value)).strip("_")
        return cleaned or "pliv_batch"

    def _batch_base_token(self, complex_name: str, ligands: List[LigandTarget]) -> str:
        if len(ligands) == 1:
            return self._sanitize_token(ligands[0].label)
        complex_token = self._sanitize_token(complex_name)
        return f"{complex_token}_{len(ligands)}ligands"

    def _session_metadata_path(self, session_path: Path) -> Path:
        return Path(str(session_path) + ".pliv.json")
    def _capture_analysis_state(self) -> dict[str, Any]:
        if self._last_request is None:
            return {}

        payload: dict[str, Any] = {
            "request": self._serialize_analysis_request(self._last_request),
        }
        if self._last_summary is not None:
            payload["summary"] = self._serialize_analysis_summary(self._last_summary)
        return payload

    def _restore_analysis_state(self, payload: Any) -> bool:
        if not isinstance(payload, dict):
            return False

        request = self._deserialize_analysis_request(payload.get("request"))
        if request is None:
            self._last_request = None
            self._last_summary = None
            self._publication_active = False
            return False

        self._last_request = request
        self._last_summary = self._deserialize_analysis_summary(payload.get("summary"))
        self._publication_active = False
        self.dialog.set_selection_mode(request.input_mode)
        self.dialog.select_complex(request.complex_name)
        self.dialog.select_ligand(request.ligand)
        self.dialog.set_interaction_types(request.interaction_types)
        return True

    def _serialize_analysis_request(self, request: AnalysisRequest) -> dict[str, Any]:
        return asdict(request)

    def _deserialize_analysis_request(self, payload: Any) -> Optional[AnalysisRequest]:
        if not isinstance(payload, dict):
            return None

        ligand_payload = payload.get("ligand")
        if not isinstance(ligand_payload, dict):
            return None

        try:
            ligand = LigandTarget(
                chain=str(ligand_payload.get("chain", "")),
                resi=str(ligand_payload.get("resi", "")),
                resn=str(ligand_payload.get("resn", "")),
                object_name=str(ligand_payload.get("object_name", "")),
            )
            return AnalysisRequest(
                input_mode=str(payload.get("input_mode", "complex")),
                complex_name=str(payload["complex_name"]),
                ligand=ligand,
                interaction_types=[str(item) for item in payload.get("interaction_types", [])],
                profile_name=str(payload.get("profile_name", self.config.default_profile())),
                token=str(payload["token"]),
                ligand_selection=str(payload["ligand_selection"]),
                protein_selection=str(payload["protein_selection"]),
                pocket_selection=str(payload["pocket_selection"]),
                water_selection=str(payload["water_selection"]),
                ion_selection=str(payload["ion_selection"]),
                group_name=str(payload["group_name"]),
            )
        except Exception:
            return None

    def _serialize_analysis_summary(self, summary: Any) -> dict[str, Any]:
        payload = {
            "ligand_atom_count": int(getattr(summary, "ligand_atom_count", 0)),
            "pocket_atom_count": int(getattr(summary, "pocket_atom_count", 0)),
            "selected_interaction_types": list(getattr(summary, "selected_interaction_types", [])),
            "notes": list(getattr(summary, "notes", [])),
        }

        for key in ("interactions", "interaction_counts", "water_atom_count", "ion_atom_count"):
            if hasattr(summary, key):
                payload[key] = getattr(summary, key)
        return payload

    def _deserialize_analysis_summary(self, payload: Any) -> Optional[AnalysisSummary]:
        if not isinstance(payload, dict):
            return None

        summary = AnalysisSummary(
            ligand_atom_count=int(payload.get("ligand_atom_count", 0)),
            pocket_atom_count=int(payload.get("pocket_atom_count", 0)),
            selected_interaction_types=[str(item) for item in payload.get("selected_interaction_types", [])],
            notes=[str(item) for item in payload.get("notes", [])],
        )
        for key in ("interactions", "interaction_counts", "water_atom_count", "ion_atom_count"):
            if key in payload:
                setattr(summary, key, payload[key])
        return summary

    def _capture_pymol_window_state(self) -> dict:
        state = {}
        window = self._find_pymol_main_window()
        if window is not None:
            state["window_size"] = {"width": int(window.width()), "height": int(window.height())}
            state["window_position"] = {"x": int(window.x()), "y": int(window.y())}
        try:
            width, height = self.engine.cmd.get_viewport()
            state["viewport"] = {"width": int(width), "height": int(height)}
        except Exception:
            pass
        return state

    def _restore_pymol_window_state(self, state: dict) -> None:
        if not state:
            self.dialog.log("No saved window metadata was available.")
            return

        window = self._find_pymol_main_window()
        if window is not None:
            size = state.get("window_size", {})
            if "width" in size and "height" in size:
                window.resize(int(size["width"]), int(size["height"]))
            position = state.get("window_position", {})
            if "x" in position and "y" in position:
                window.move(int(position["x"]), int(position["y"]))

        viewport = state.get("viewport", {})
        if "width" in viewport and "height" in viewport:
            try:
                self.engine.cmd.viewport(int(viewport["width"]), int(viewport["height"]))
            except Exception:
                pass

        self.dialog.log("Restored saved PyMOL window size and viewport.")

    def _find_pymol_main_window(self):
        app = QtWidgets.QApplication.instance()
        if app is None:
            return None
        for widget in app.topLevelWidgets():
            title = widget.windowTitle().strip()
            if title == "PyMOL" or title.startswith("PyMOL "):
                return widget
        return None

    def store_viewpoint(self, slot: int) -> None:
        view_name = self.renderer.store_viewpoint(slot)
        self.dialog.log(f"Stored viewpoint {slot} as {view_name}.")

    def restore_viewpoint(self, slot: int) -> None:
        restored = self.renderer.restore_viewpoint(slot)
        if restored:
            self.dialog.log(f"Restored viewpoint {slot}.")
        else:
            self.dialog.log(f"No saved viewpoint was available in slot {slot}.")
